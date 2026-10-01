"""
The two review workbooks for the open-data build (docs/release/06-OUR-OWN-DATA.md):

  docs/release/open-data/La Liga ratings, all seasons.xlsx   every player of every
      season, the top 50 of the last complete one, the clubs, and how each rating
      is made (its four parts are columns, so any row can be checked by hand)
  docs/release/open-data/Open data coverage.xlsx         what the rebuild kept and
      lost against today's database: per league-season, per club, and the clubs
      Wikipedia had that the game doesn't

  python scripts/open-data-sheets.py      (after build-open-seeds.ts)

Today's (Transfermarkt-built) ratings appear as a yardstick column only.
"""
import json
import sqlite3
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / "scripts" / "seed-open" / "_report.json"
OUT = ROOT / "docs" / "release" / "open-data"
# The old (Transfermarkt-built) database, only if given as TM_DB: the bundled
# one is the open data now, so the "today" columns would compare it with itself.
import os
DB = Path(os.environ["TM_DB"]) if os.environ.get("TM_DB") else None

FONT = "Arial"
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
BODY = Font(name=FONT)
NOTE = Font(name=FONT, italic=True, color="555555")


def norm(s):
    s = unicodedata.normalize("NFD", s or "")
    return " ".join("".join(c for c in s if not unicodedata.combining(c)).lower().replace("-", " ").split())


def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def body(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font != NOTE:
                c.font = BODY


def la_liga(report):
    seasons = report["ratings"].get("la_liga")
    if not seasons:
        print("no La Liga ratings in the report; run the build first")
        return
    # Yardstick: today's La Liga, by season and name.
    today = {}
    con = sqlite3.connect(DB) if DB else None
    for n, o, y in (con.execute(
            "SELECT p.name, ps.ovr, cs.year_start FROM player_seasons ps JOIN players p ON p.id = ps.player_id "
            "JOIN club_seasons cs ON cs.id = ps.club_season_id JOIN clubs c ON c.id = cs.club_id "
            "WHERE c.league_id = 'la_liga'") if con else []):
        today[(y, norm(n))] = o
    if con:
        con.close()

    # PLAIN VALUES, not formulas: formulas carried no stored results (nothing on
    # this machine can compute them), so previews and some viewers showed blank
    # columns. The arithmetic is spelled out on the "How it's rated" sheet.
    cols = ["Season", "Player", "Club", "Club finish", "Position", "Nationality", "Born", "Age",
            "League apps", "League goals", "All comps apps", "All comps goals", "Games he could play",
            "Share of games", "Goals per game",
            "Club level", "Role", "Production", "Age adj.", "OVR",
            "Attack", "Defence", "Physical", "Pace", "Technical",
            "Today's OVR (yardstick)", "Ours − today", "Data", "Wikipedia article"]
    widths = [9, 24, 22, 8, 8, 13, 7, 6, 8, 8, 9, 9, 9, 8, 8, 8, 7, 9, 7, 6, 7, 7, 8, 6, 9, 11, 9, 8, 32]
    pct_cols, dp2_cols = {14}, {15, 16, 17, 18, 19}

    def row_of(year, p):
        age = year - p["birthYear"] if p.get("birthYear") else None
        apps, goals, games = p.get("apps"), p.get("goals"), p["games"]
        t = today.get((year, norm(p["name"])))
        parts, a = p["parts"], p["attrs"]
        return [
            f"{year}-{(year + 1) % 100:02d}", p["name"], p["club"], p["clubPosition"], p.get("position") or "",
            p.get("nationality") or "", p.get("birthYear"), age,
            apps, goals, p.get("totalApps"), p.get("totalGoals"), games,
            (apps / games) if apps is not None and games else None,
            (goals / apps) if goals is not None and apps else None,
            parts["club"], parts["role"], parts["production"], parts["age"], parts["ovr"],
            a["attack"], a["defense"], a["physical"], a["pace"], a["technical"],
            t, (parts["ovr"] - t) if t is not None else None, p["level"], p.get("article") or "",
        ]

    def fill(ws, rows):
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.column in pct_cols:
                    c.number_format = "0%"
                elif c.column in dp2_cols:
                    c.number_format = "0.00"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
        body(ws)

    wb = Workbook()
    wb.remove(wb.active)
    years = sorted(int(y) for y in seasons)
    latest_done = max(y for y in years if y < 2025)  # the last COMPLETE season
    latest = seasons[str(latest_done)]
    fill(sheet(wb, f"Top 50 {latest_done}-{(latest_done + 1) % 100:02d}", cols, widths),
         [row_of(latest_done, p) for p in sorted(latest, key=lambda p: -p["parts"]["ovr"])[:50]])
    everyone = [row_of(y, p) for y in years for p in sorted(seasons[str(y)], key=lambda p: (p["clubPosition"], -p["parts"]["ovr"]))]
    fill(sheet(wb, "All players, all seasons", cols, widths), everyone)

    # Per club-season summary, computed here (values, same reason as above).
    ws = sheet(wb, "Clubs by season", ["Season", "Club", "Finish", "Players", "Average OVR", "Best OVR", "Team strength (top 14)", "Today's average (yardstick)"], [9, 26, 8, 8, 10, 9, 12, 13])
    for y in years:
        by = {}
        for p in seasons[str(y)]:
            by.setdefault((p["clubPosition"], p["club"]), []).append(p)
        for (pos, club), ps in sorted(by.items()):
            o = sorted((p["parts"]["ovr"] for p in ps), reverse=True)
            t = [today[(y, norm(p["name"]))] for p in ps if (y, norm(p["name"])) in today]
            ws.append([f"{y}-{(y + 1) % 100:02d}", club, pos, len(ps), sum(o) / len(o), o[0],
                       sum(o[:14]) / len(o[:14]), (sum(t) / len(t)) if t else None])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column in (5, 7, 8):
                c.number_format = "0.0"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"
    body(ws)

    ws = wb.create_sheet("How it's rated")
    ws.column_dimensions["A"].width = 120
    lines = [
        "How the open-data OVR is made (scripts/lib/open-rating.ts)",
        "",
        "OVR = club level + role + production + age adj., clamped to the game's scale (58-93) and rounded. The four parts are in the columns before OVR; add them up to check any row.",
        "Club level: La Liga's band from Spain's UEFA association rank (3): champion 85.7, last place 77.4, on a curve (strength^1.3) so the top of the table pulls away.",
        "Role: share of the league games he could have played (League apps / Games he could play). A regular (55%) sits at his club's level; above that +8 per 100%, below it -14 per 100%.",
        "  A January signing's 'games he could play' is half the season. 2025-26 is in progress: games so far, from the league's most-used player.",
        "Production: goals per game against what his position usually scores (ST 0.40, winger 0.25, CAM 0.20, CM 0.10, full-back 0.04, CB 0.05); needs 8+ apps; never keepers.",
        "  Weighted by position (striker 1, winger 0.85, CAM 0.8, wide mid 0.7, CM 0.55, CDM 0.35, full-back 0.3, centre-back 0.25), capped at -1.5 to +2.5 times that weight.",
        "Age adj.: 22 and under lose 0.9 per year below 23, shrinking with playing time (x (1 - 0.8 x share)); 33 and over lose 0.8 per year above 32.",
        "Attack / Defence / Physical / Pace / Technical: the game's own per-position spread around the OVR (scripts/lib/transfermarkt.ts attributes(); PoM's code, no outside data).",
        "",
        "Stats: League apps and goals are the rating's inputs. All-competitions apps and goals (league + cups + Europe) come from the same career-table row, shown for judging; they don't feed the rating yet.",
        "Assists and minutes: not available. Wikipedia's career tables record appearances and goals only, so the rebuild has no assists or minutes; the game shows 0 assists until an open source is found.",
        "Data: 'full' = a career-table row for this club and season. 'partial' = his page has no career table; league apps estimated from his infobox spell (apps at the club over the spell's seasons), no goals.",
        "Who is listed: players with at least one league appearance for the club that season (today's database also lists unused squad members, so its squads are bigger).",
        "Today's OVR: the current, Transfermarkt-built rating for the same player-season, matched by name. A yardstick for judging, never an input: no constant was fitted to it.",
    ]
    for i, t in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=(i == 1), size=12 if i == 1 else 10)
        c.alignment = Alignment(wrap_text=True)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "La Liga ratings, all seasons.xlsx"
    wb.save(path)
    print(f"wrote {path} ({len(everyone)} player-seasons over {len(years)} seasons)")


def coverage(report):
    rows = report["coverage"]
    if not rows:
        return
    wb = Workbook()
    wb.remove(wb.active)
    n = len(rows) + 1

    # Values, not formulas: nothing here can compute them, and viewers showed
    # formula cells blank.
    ws = sheet(wb, "By league-season", ["League", "Season", "Clubs today", "Clubs rebuilt", "Players today", "Players rebuilt", "Players kept"], [26, 8, 11, 11, 12, 13, 11])
    groups = {}
    for r in rows:
        g = groups.setdefault((r["league"], r["season"]), [0, 0, 0, 0])
        g[0] += 1; g[1] += 1 if r["after"] > 0 else 0; g[2] += r["before"]; g[3] += r["after"]
    tot = [0, 0, 0, 0]
    for (lg, season), g in sorted(groups.items()):
        ws.append([lg, season, g[0], g[1], g[2], g[3], (g[3] / g[2]) if g[2] else None])
        ws.cell(row=ws.max_row, column=7).number_format = "0%"
        tot = [a + b for a, b in zip(tot, g)]
    ws.append([])
    ws.append(["All", "", tot[0], tot[1], tot[2], tot[3], (tot[3] / tot[2]) if tot[2] else None])
    ws.cell(row=ws.max_row, column=7).number_format = "0%"
    for c in ws[ws.max_row]:
        c.font = Font(name=FONT, bold=True)
    body(ws)

    ws = sheet(wb, "By club", ["League", "Season", "Club (game)", "Players today", "Players rebuilt", "Kept", "Wikipedia article", "Full data", "Partial", "Bare", "Note"], [26, 8, 30, 12, 13, 8, 34, 9, 8, 8, 30])
    for i, r in enumerate(sorted(rows, key=lambda r: (r["league"], r["season"], r["club"])), start=2):
        ws.append([r["league"], r["season"], r["club"], r["before"], r["after"], (r["after"] / r["before"]) if r["before"] else None,
                   r.get("wikipedia") or "(not found)", r.get("full"), r.get("partial"), r.get("bare"), r.get("note")])
        ws.cell(row=i, column=6).number_format = "0%"
    ws.auto_filter.ref = f"A1:K{ws.max_row}"
    body(ws)

    ws = sheet(wb, "Wikipedia only", ["League", "Season", "Wikipedia club (no match in the game)", "Players"], [26, 8, 40, 9])
    for r in report["unmatched"]:
        ws.append([r["league"], r["season"], r["wikipedia"], r["players"]])
    body(ws)

    wb.calculation.fullCalcOnLoad = True
    ws = sheet(wb, "What was patched", ["What", "League", "Club", "Season", "Detail"], [26, 22, 30, 8, 60])
    for t in report.get("toppedUp", []):
        ws.append(["Squad topped up", "", t["club"], 2025, f"{t['added']} players added from the club's {'own 2025-26 history' if t.get('from') == 'history' else 'Wikidata squad records (rated as low-data players)'}"])
    for t in report.get("proxied", []):
        ws.append(["Borrowed current squad", "Champions League", t["club"], t["season"], f"only {t['found']} players found for that season; its 2025-26 squad stands in"])
    for t in report.get("dropped", []):
        ws.append(["Left out of the pool", t["league"], t["club"], 2025, f"{t['players']} players even after top-ups; finished {t['position']}th, so no European place moves"])
    body(ws)

    path = OUT / "Open data coverage.xlsx"
    wb.save(path)
    print(f"wrote {path} ({len(rows)} club-seasons)")


if __name__ == "__main__":
    import sys
    # An optional path: a trial build's report (OPEN_OUT) instead of the real one.
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else REPORT
    rep = json.loads(src.read_text(encoding="utf-8"))
    la_liga(rep)
    coverage(rep)
